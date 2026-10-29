"""
Pruebas sin internet: simulan la API de iNaturalist.
Correr con:  pytest -q
"""
import io
import json

import openpyxl
import pytest
from PIL import Image

import sepia
from sepia import inat, config

LAT, LON = 57.1900283333, -2.072345


def fake_obs(i):
    return {"id": i, "uuid": f"u{i}", "observed_on": "2026-02-22", "quality_grade": "research",
            "captive": False, "license_code": "cc0", "place_guess": "Bridge of Don, UK",
            "geojson": {"type": "Point", "coordinates": [LON, LAT]},
            "positional_accuracy": 10, "obscured": i == 3, "user": {"login": "x"},
            "taxon": {"id": 151429, "name": "Sepia officinalis"},
            "observation_photos": [
                {"position": 1, "photo": {"id": i * 10 + 2, "license_code": "cc0",
                 "url": f"https://inaturalist-open-data.s3.amazonaws.com/photos/{i*10+2}/square.jpg"}},
                {"position": 0, "photo": {"id": i * 10 + 1, "license_code": None,
                 "url": f"https://static.inaturalist.org/photos/{i*10+1}/square.jpeg?123"}}]}


ALL = [fake_obs(i) for i in range(1, 8)]


class R:
    def __init__(self, j=None, c=b""):
        self._j, self.content, self.status_code = j, c, 200
    def json(self):
        return self._j
    def raise_for_status(self):
        pass


@pytest.fixture
def api(monkeypatch):
    buf = io.BytesIO()
    Image.new("RGB", (2048, 1536), (200, 200, 200)).save(buf, "JPEG")
    jpg = buf.getvalue()

    def fake_get(self, url, params=None, timeout=None):
        if url == config.API_URL:
            if params.get("per_page") == 0:
                return R({"total_results": len(ALL), "results": []})
            if "id" in params:
                ids = {int(x) for x in params["id"].split(",")}
                return R({"results": [o for o in ALL if o["id"] in ids]})
            return R({"results": [o for o in ALL if o["id"] > params["id_above"]][:3]})
        assert "/original." in url
        return R(c=jpg)

    monkeypatch.setattr(inat.requests.Session, "get", fake_get)
    monkeypatch.setattr(inat, "API_PAUSE_S", 0)
    monkeypatch.setattr(inat, "PHOTO_PAUSE_S", 0)


@pytest.fixture
def cfg(tmp_path):
    return sepia.Config(base_dir=str(tmp_path), per_page=3)


def test_download_resume_and_tables(api, cfg):
    sepia.download_observations(cfg, max_pages=2)   # se corta a las 2 paginas
    sepia.download_observations(cfg)                # reanuda
    assert len(sepia.load_raw(cfg)) == 7
    obs, pho = sepia.build_tables(cfg)
    assert list(pho[pho.observation_id == 1].photo_id) == [11, 12]   # orden del observador
    assert pho.file_name[0] == "1_11.jpg"
    assert pho.download_url[0].endswith("/original.jpeg?123")
    back, _ = sepia.load_tables(cfg)
    assert back.latitude[0] == LAT and back.longitude[0] == LON       # precision completa


def test_photos_resume_and_resize(api, cfg):
    sepia.download_observations(cfg)
    sepia.build_tables(cfg)
    sepia.download_photos(cfg, limit=5)
    stats = sepia.download_photos(cfg)
    assert stats["skipped"] == 5 and stats["ok"] == 9
    sepia.make_resized(cfg, 512)
    small = next((sepia.paths(cfg)["base"] / "photos" / "512px").glob("*.jpg"))
    assert max(Image.open(small).size) == 512
    rep = sepia.check_download(cfg)
    assert rep["photos_missing"] == 0 and rep["coordinates_obscured"] == 1


def test_download_by_ids(api, cfg):
    assert sepia.download_by_ids(cfg, [1, 2, 99, 5]) == 3


def test_excel(api, cfg):
    sepia.download_observations(cfg)
    sepia.build_tables(cfg)
    ws = openpyxl.load_workbook(sepia.build_excel(cfg, "a.xlsx"))["identification"]
    assert ws.max_row == 1 + 14 and ws.cell(2, 4).value == LAT     # 7 registros x 2 fotos
    assert not any(isinstance(c.value, str) and c.value.startswith("=")
                   for row in ws.iter_rows() for c in row)

    # Registro 1: fotos 11 (dorsal) y 12 (ventral) del individuo a; en la foto 12
    # aparece ademas un segundo hueso b. El lado es atributo de la fila.
    f = sepia.paths(cfg)["annotations"] / "HZ.jsonl"
    sepia.append_jsonl({"observation_id": 1, "photo_id": 11, "individual": "a", "side": "dorsal",
                        "epibionts": "yes", "saved_at": "2026-10-28T10:00"}, f)
    sepia.append_jsonl({"observation_id": 1, "photo_id": 12, "individual": "a", "side": "ventral",
                        "bitemarks": "no", "saved_at": "2026-10-28T10:01"}, f)
    sepia.append_jsonl({"observation_id": 1, "photo_id": 12, "individual": "a", "side": "ventral",
                        "bitemarks": "yes", "saved_at": "2026-10-28T10:05"}, f)   # correccion
    sepia.append_jsonl({"observation_id": 1, "photo_id": 12, "individual": "b", "side": "dorsal",
                        "saved_at": "2026-10-28T10:06"}, f)
    sepia.append_jsonl({"observation_id": 1, "individual": "c", "side": "dorsal",
                        "saved_at": "2026-10-28T10:07"}, f)                       # sin photo_id: se ignora
    wb = openpyxl.load_workbook(sepia.build_excel(cfg, "b.xlsx"))
    ws = wb["identification"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    r1 = [r for r in rows if r["ID"] == 1]
    assert len(rows) == 15 and len(r1) == 3
    assert [(r["photo_id"], r["individual (a,b,c,d,e)"]) for r in r1] == [(11, "a"), (12, "a"), (12, "b")]
    ventral = [r for r in r1 if r["ventral side (cuttlebone)"] == 1]
    assert len(ventral) == 1 and ventral[0]["¿bitemarks?"] == "yes"   # vale la ultima
    assert all(r["number individuals (cuttlebone)"] == 2 for r in r1)  # letras distintas, no filas
    assert all(r["latitude"] == LAT for r in r1)                       # coordenadas por join
    assert wb.sheetnames == ["identification", "legend", "excluded"]


def test_config_yaml(tmp_path):
    y = tmp_path / "c.yaml"
    y.write_text("base_dir: /x\nbbox: [48, -6, 59.5, 13]\nresize: [512]\n")
    c = sepia.load_config(str(y))
    assert c.bbox == (48, -6, 59.5, 13) and c.resize == (512,)
    assert c.query_params()["swlat"] == 48


def test_ids_file(tmp_path):
    f = tmp_path / "ids.txt"
    f.write_text("# comentario\n339768773\n\n26739989\n339768773\n")
    assert sepia.ids_from_file(str(f)) == [26739989, 339768773]


def test_pilot_ids_in_repo():
    ids = sepia.ids_from_file("configs/pilot79_ids.txt")
    assert len(ids) == 79


class FakeProvider:
    name = "fake"
    def __init__(self):
        self.calls = []
    def classify(self, images, prompt):
        self.calls.append(len(images))
        res = {"photos": [{"photo_index": i + 1, "content": "cuttlebone", "n_cuttlebones_visible": 1}
                          for i in range(len(images))],
               "record_state": "cuttlebone", "n_cuttlebones": 1, "count_possible": True,
               "same_bone_in_several_photos": len(images) > 1, "confidence": "high", "notes": ""}
        sepia.triage.TriageResult.model_validate(res)   # cumple el esquema
        return res, json.dumps(res), {"input_tokens": 500, "output_tokens": 80}


def test_triage_resume(api, cfg, monkeypatch):
    sepia.download_observations(cfg)
    sepia.build_tables(cfg)
    sepia.download_photos(cfg)
    sepia.make_resized(cfg, 512)
    cfg.triage = {"provider": "fake", "model": "m1", "requests_per_minute": 100000}
    monkeypatch.setattr(sepia.triage.time, "sleep", lambda s: None)
    prov = FakeProvider()
    sepia.run_triage(cfg, provider=prov, limit=3)
    sepia.run_triage(cfg, provider=prov)              # reanuda: solo los 4 que faltan
    assert len(prov.calls) == 7 and set(prov.calls) == {2}   # todas las fotos del registro juntas
    t = sepia.load_triage(cfg, "fake", "m1")
    assert len(t) == 7 and t["n_cuttlebones"].eq(1).all()
    raw = sepia.read_jsonl(sepia.triage.output_path(cfg, "fake", "m1"))
    assert raw["prompt_sha"].nunique() == 1 and raw["photo_ids"].iloc[0] == [11, 12]


def test_prompt_versioned():
    text, sha = sepia.load_prompt("triage_v1")
    assert "DISTINCT cuttlebones" in text and len(sha) == 12


def test_triage_stops_on_fatal(api, cfg, monkeypatch):
    sepia.download_observations(cfg)
    sepia.build_tables(cfg)
    sepia.download_photos(cfg)
    sepia.make_resized(cfg, 512)
    cfg.triage = {"provider": "fake", "model": "gone", "requests_per_minute": 100000}
    monkeypatch.setattr(sepia.triage.time, "sleep", lambda s: None)

    class Gone:
        calls = 0
        def classify(self, images, prompt):
            Gone.calls += 1
            raise sepia.triage.FatalProviderError("404 NOT_FOUND model gone")

    sepia.run_triage(cfg, provider=Gone())
    assert Gone.calls == 1                       # se detuvo al primer error de configuracion
    assert sepia.triage.done_ids(sepia.triage.output_path(cfg, "fake", "gone")) == set()


def test_triage_retries_overload(api, cfg, monkeypatch):
    sepia.download_observations(cfg)
    sepia.build_tables(cfg)
    sepia.download_photos(cfg)
    sepia.make_resized(cfg, 512)
    cfg.triage = {"provider": "fake", "model": "busy", "requests_per_minute": 100000}
    monkeypatch.setattr(sepia.triage.time, "sleep", lambda s: None)
    ok = FakeProvider()

    class Busy:
        n = 0
        def classify(self, images, prompt):
            Busy.n += 1
            if Busy.n == 1:   # primer intento: modelo saturado
                raise sepia.triage.RateLimited("503 UNAVAILABLE high demand")
            return ok.classify(images, prompt)

    sepia.run_triage(cfg, provider=Busy(), limit=2)
    t = sepia.load_triage(cfg, "fake", "busy")
    assert len(t) == 2 and Busy.n == 3   # reintento y luego siguio


def test_triage_only_ids(api, cfg, monkeypatch):
    sepia.download_observations(cfg)
    sepia.build_tables(cfg)
    sepia.download_photos(cfg)
    sepia.make_resized(cfg, 512)
    cfg.triage = {"provider": "fake", "model": "ids", "requests_per_minute": 100000}
    monkeypatch.setattr(sepia.triage.time, "sleep", lambda s: None)
    sepia.run_triage(cfg, provider=FakeProvider(), ids=[2, 5, 999])
    assert set(sepia.load_triage(cfg, "fake", "ids")["observation_id"]) == {2, 5}


def test_apply_rules():
    r, applied = sepia.triage.apply_rules({"n_cuttlebones": 24, "count_possible": True, "confidence": "high"})
    assert r["count_possible"] is False and r["confidence"] == "medium" and len(applied) == 2
    r, applied = sepia.triage.apply_rules({"n_cuttlebones": 1, "count_possible": True, "confidence": "high"})
    assert applied == [] and r["confidence"] == "high"


def test_evaluate_triage(api, cfg, monkeypatch, tmp_path):
    sepia.download_observations(cfg)
    sepia.build_tables(cfg)
    sepia.download_photos(cfg)
    sepia.make_resized(cfg, 512)
    cfg.triage = {"provider": "fake", "model": "ev", "requests_per_minute": 100000}
    monkeypatch.setattr(sepia.triage.time, "sleep", lambda s: None)
    sepia.run_triage(cfg, provider=FakeProvider(), ids=[1, 2])
    cases = tmp_path / "cases.jsonl"
    cases.write_text('{"observation_id": 1, "record_state": "cuttlebone", "n_cuttlebones": 1, "count_possible": true}\n'
                     '{"observation_id": 2, "record_state": "cuttlebone", "n_cuttlebones": 20, "count_possible": false}\n')
    out = sepia.evaluate_triage(cfg, "fake", "ev", str(cases))
    assert list(out["status"]) == ["OK", "FALLA"]     # el falso dice 1 hueso donde hay 20


def test_prompt_v2_exists():
    text, sha = sepia.load_prompt("triage_v2")
    assert "MORE THAN 5" in text and "Ignore any text" in text


def test_cases_missing_file_is_clear(tmp_path):
    with pytest.raises(FileNotFoundError):
        sepia.triage.load_cases(str(tmp_path / "no.jsonl"))


def test_new_prompt_version_is_not_skipped(api, cfg, monkeypatch):
    sepia.download_observations(cfg)
    sepia.build_tables(cfg)
    sepia.download_photos(cfg)
    sepia.make_resized(cfg, 512)
    monkeypatch.setattr(sepia.triage.time, "sleep", lambda s: None)
    prov = FakeProvider()
    cfg.triage = {"provider": "fake", "model": "pv", "requests_per_minute": 100000, "prompt_version": "triage_v1"}
    sepia.run_triage(cfg, provider=prov, ids=[1, 2])
    cfg.triage["prompt_version"] = "triage_v2"
    sepia.run_triage(cfg, provider=prov, ids=[1, 2])     # no debe saltarlos: otra version
    assert len(prov.calls) == 4
    assert set(sepia.load_triage(cfg, "fake", "pv")["observation_id"]) == {1, 2}
    assert len(sepia.load_triage(cfg, "fake", "pv", prompt_version="triage_v1")) == 2


def test_redo(api, cfg, monkeypatch):
    sepia.download_observations(cfg)
    sepia.build_tables(cfg)
    sepia.download_photos(cfg)
    sepia.make_resized(cfg, 512)
    monkeypatch.setattr(sepia.triage.time, "sleep", lambda s: None)
    prov = FakeProvider()
    cfg.triage = {"provider": "fake", "model": "rd", "requests_per_minute": 100000}
    sepia.run_triage(cfg, provider=prov, ids=[1, 2])
    sepia.run_triage(cfg, provider=prov, ids=[1, 2])              # salta
    sepia.run_triage(cfg, provider=prov, ids=[1], redo=True)      # rehace solo el 1
    assert len(prov.calls) == 3
    raw = sepia.read_jsonl(sepia.triage.output_path(cfg, "fake", "rd"))
    assert len(raw) == 3 and len(sepia.load_triage(cfg, "fake", "rd")) == 2


def test_excel_with_triage(api, cfg, monkeypatch):
    """El triaje llena state y n_cuttlebones; una anotacion manual tiene prioridad."""
    sepia.download_observations(cfg)
    sepia.build_tables(cfg)
    sepia.download_photos(cfg)
    sepia.make_resized(cfg, 512)
    cfg.triage = {"provider": "fake", "model": "m1", "prompt_version": "triage_v2",
                  "requests_per_minute": 100000}
    monkeypatch.setattr(sepia.triage.time, "sleep", lambda s: None)
    sepia.run_triage(cfg, provider=FakeProvider())
    f = sepia.paths(cfg)["annotations"] / "HZ.jsonl"
    sepia.append_jsonl({"observation_id": 2, "photo_id": 21, "individual": "a", "side": "dorsal",
                        "state": "egg case", "label_source": "annotator",
                        "saved_at": "2026-10-28T10:00"}, f)
    ws = openpyxl.load_workbook(sepia.build_excel(cfg, "t.xlsx"))["identification"]
    hdr = [c.value for c in ws[1]]
    rows = {(r[1], r[hdr.index("photo_id")]): dict(zip(hdr, r))
            for r in ws.iter_rows(min_row=2, values_only=True)}
    r1, r2 = rows[(1, 11)], rows[(2, 21)]
    assert r1["state"] == "cuttlebone" and r1["number individuals (cuttlebone)"] == 1
    assert r1["label_source"] == "gpt" and r1["model_version"] == "m1/triage_v2"
    assert json.loads(r1["model_prob"])["confidence"] == "high"
    assert r1["reviewed"] is False
    assert r2["state"] == "egg case" and r2["label_source"] == "annotator"   # gana la persona


# --------------------------------------------------------------------------
# Anotacion manual
# --------------------------------------------------------------------------

def _ready(cfg):
    sepia.download_observations(cfg)
    sepia.build_tables(cfg)
    sepia.download_photos(cfg)
    sepia.make_resized(cfg, 1024)


def test_annotation_events(api, cfg):
    """Dorsal y ventral del mismo individuo en fotos distintas conviven; corregir = la ultima."""
    from sepia import annotate as an
    _ready(cfg)
    q = an.work_queue(cfg)
    assert len(q) == 14 and an.first_pending(q, an.load_events(cfg)) == 0
    an.save_individual(cfg, "HZ", 1, 11, "a", {"state": "cuttlebone", "side": "dorsal",
                                               "notes": "same bone as in record 2?"})
    an.save_individual(cfg, "HZ", 1, 12, "a", {"state": "cuttlebone", "side": "dorsal"})
    an.save_individual(cfg, "HZ", 1, 12, "a", {"state": "cuttlebone", "side": "ventral"})  # correccion
    an.save_individual(cfg, "HZ", 1, 12, "b", {"state": "cuttlebone", "side": "dorsal"})
    an.delete_individual(cfg, "HZ", 1, 12, "b")                                           # error
    an.mark_empty_photo(cfg, "HZ", 2, 21)
    an.exclude_record(cfg, "HZ", 3, "captive animal")
    an.exclude_record(cfg, "HZ", 4, "duplicate")
    an.exclude_record(cfg, "HZ", 4, "")                                                   # deshacer
    ev = an.load_events(cfg)
    ind = ev["individual"].set_index(["photo_id", "individual"])["side"].to_dict()
    assert ind == {(11, "a"): "dorsal", (12, "a"): "ventral"}
    assert an.letters_in_record(ev, 1) == ["a"] and an.next_letter(["a", "b"]) == "c"
    assert list(ev["exclude_record"]["observation_id"]) == [3]
    # pendiente: foto 22 (la 21 esta vacia, el registro 1 completo)
    assert q.iloc[an.first_pending(q, ev)]["photo_id"] == 22
    line = sepia.read_jsonl(an.annotation_file(cfg, "HZ")).iloc[0]
    assert line["schema_version"] == 1 and line["annotator"] == "HZ" and line["saved_at"].endswith("+00:00")

    wb = openpyxl.load_workbook(sepia.build_excel(cfg, "ann.xlsx"))
    hdr = [c.value for c in wb["identification"][1]]
    rows = [dict(zip(hdr, r)) for r in wb["identification"].iter_rows(min_row=2, values_only=True)]
    r1 = [r for r in rows if r["ID"] == 1]
    assert [(r["photo_id"], r["dorsal side (cuttlebone)"], r["ventral side (cuttlebone)"]) for r in r1] \
        == [(11, 1, None), (12, None, 1)]
    assert r1[0]["notes"] == "same bone as in record 2?"
    assert all(r["number individuals (cuttlebone)"] == 1 and r["reviewed"] is True for r in r1)
    r21 = [r for r in rows if r["photo_id"] == 21][0]
    assert r21["reviewed"] is True and r21["individual (a,b,c,d,e)"] is None
    assert not any(r["ID"] == 3 for r in rows)
    ex = list(wb["excluded"].iter_rows(min_row=2, values_only=True))
    assert len(ex) == 1 and ex[0][0] == 3 and ex[0][-1] == "captive animal"


def test_annotation_app_smoke(api, cfg):
    """La interfaz se construye y los botones escriben las lineas esperadas."""
    pytest.importorskip("ipywidgets")
    _ready(cfg)
    app = sepia.AnnotationApp(cfg, annotator="HZ", round_=1)
    assert app.row["photo_id"] == 11 and len(app.image.value) > 0
    assert len(app.strip.children) == 2                           # miniaturas de las 2 fotos
    app.strip.children[1].children[1].click()                     # clic en la miniatura 2
    assert app.row["photo_id"] == 12
    app.strip.children[0].children[1].click()
    assert app.row["photo_id"] == 11
    app.inputs["state"].value = "cuttlebone"
    assert app.inputs["bitemarks"].value == "no" and app.inputs["epibionts"].value == "no"  # sugeridos
    assert app.inputs["side"].layout.display is None              # visible para huesos
    assert app.inputs["barnacles_n"].layout.display == "none"     # oculto sin epibiontes
    app.inputs["side"].value = "dorsal"
    app.inputs["epibionts"].value = "yes"
    app.inputs["barnacles_n"].value = "1-10"
    app.buttons.children[0].click()                               # "Save ▶": guarda y avanza
    assert "Saved" in app.msg.value and app.row["photo_id"] == 12
    assert app.letter.value == "b" and "a" in [o[1] for o in app.letter.options]
    app.letter.value = "a"                                        # mismo hueso, otra foto
    assert app.inputs["state"].value == "cuttlebone" and app.inputs["side"].value is None
    assert app.inputs["epibionts"].value == "no"                  # sugerido (nuevo en esta foto)
    app.inputs["side"].value = "ventral"
    app.inputs["epibionts"].value = "no"
    app._safe(app.on_save)
    app._safe(app.on_next)
    app._safe(app.on_empty)                                       # foto 21 sin individuos
    app.reason.value = "not Sepia"
    app._safe(app.on_exclude)                                     # excluye registro 2
    app._safe(app.on_save)                                        # sin estado: error visible
    assert "Error" in app.msg.value
    ev = sepia.load_events(cfg)
    ind = ev["individual"]
    assert len(ind) == 2 and ind["round"].eq(1).all()
    assert ind.set_index("photo_id").loc[11, "barnacles_n"] == "1-10"
    assert "barnacles_n" not in ind or pd_isna(ind.set_index("photo_id").loc[12].get("barnacles_n"))
    assert list(ev["exclude_record"]["observation_id"]) == [2]
    app2 = sepia.AnnotationApp(cfg, annotator="HZ")              # retoma donde quedo
    assert app2.row["observation_id"] == 3


def pd_isna(v):
    import pandas as pd
    return v is None or pd.isna(v)


def test_saved_blank_stays_blank(api, cfg):
    """Un 'no se puede evaluar' (vacio) ya guardado no se convierte en 'no' al volver a abrirlo."""
    pytest.importorskip("ipywidgets")
    from sepia import annotate as an
    _ready(cfg)
    an.save_individual(cfg, "HZ", 1, 11, "a", {"state": "cuttlebone", "bitemarks": None,
                                               "epibionts": "yes"})
    app = sepia.AnnotationApp(cfg, annotator="HZ")
    app._go(0)
    assert app.letter.value == "a"
    assert app.inputs["bitemarks"].value is None and app.inputs["epibionts"].value == "yes"


def test_crop_box_roundtrip(tmp_path):
    """Recuadro relativo: sirve igual en 1024 px y en el original (y respeta EXIF)."""
    from sepia import annotate as an
    box = an.to_crop_box({"x": 256, "y": 192, "width": 512, "height": 384}, 1024, 768)
    assert box == {"x0": 0.25, "y0": 0.25, "x1": 0.75, "y1": 0.75}
    # dibujado de derecha a izquierda / de abajo hacia arriba, y saliendose de la foto
    assert an.to_crop_box({"x": 768, "y": 576, "width": -512, "height": -384}, 1024, 768) == box
    assert an.to_crop_box({"x": 900, "y": -50, "width": 300, "height": 100}, 1024, 768)["x1"] == 1.0
    assert an.from_crop_box(box, 2048, 1536) == {"x": 512, "y": 384, "width": 1024, "height": 768}
    f = tmp_path / "o.jpg"
    Image.new("RGB", (2048, 1536)).save(f)
    assert an.crop_photo(f, box).size == (1024, 768)


def test_crop_box_text():
    from sepia import annotate as an
    box = {"x0": 0.1, "y0": 0.2, "x1": 0.5, "y1": 0.9}
    assert an.parse_crop_box(an.format_crop_box(box)) == box
    assert an.parse_crop_box("") is None
    with pytest.raises(ValueError):
        an.parse_crop_box("0.5, 0.2, 0.1, 0.9")                  # x0 > x1


def test_annotation_app_boxes(api, cfg):
    """Boton Draw box -> arrastrar -> coordenadas en el campo -> se guardan en crop_box."""
    pytest.importorskip("jupyter_bbox_widget")
    _ready(cfg)
    app = sepia.AnnotationApp(cfg, annotator="HZ")
    assert app.bbox is not None and app.bbox.view_only is True     # no se dibuja sin el boton
    W, H = app._box_size
    app.bbox.bboxes = [{"x": 0, "y": 0, "width": W / 3, "height": H / 3, "label": "a"}]
    assert app.crop.value == ""                                    # sin boton: se ignora
    app.draw_btn.value = True
    assert app.bbox.view_only is False and app.bbox.label == "a"
    app.bbox.bboxes = [{"x": 0, "y": 0, "width": W / 2, "height": H / 2, "label": "a"}]
    assert app.crop.value == "0.0000, 0.0000, 0.5000, 0.5000"      # al soltar: el campo
    assert app.draw_btn.value is False and app.bbox.view_only is True
    app.inputs["state"].value = "cuttlebone"
    app._safe(app.on_save)
    ind = sepia.load_events(cfg)["individual"]
    assert ind.iloc[0]["crop_box"] == {"x0": 0.0, "y0": 0.0, "x1": 0.5, "y1": 0.5}
    assert len(app.bbox.bboxes) == 1 and app.crop.value.endswith("0.5000")   # recargado
    app.crop.value = "0.1, 0.1, 0.6, 0.6"                           # corregir a mano
    app._safe(app.on_save)
    assert sepia.load_events(cfg)["individual"].iloc[0]["crop_box"]["x1"] == 0.6
    app.crop.value = "nonsense"
    app._safe(app.on_save)
    assert "Error" in app.msg.value
    wb = openpyxl.load_workbook(sepia.build_excel(cfg, "box.xlsx"))["identification"]
    hdr = [c.value for c in wb[1]]
    r = dict(zip(hdr, next(wb.iter_rows(min_row=2, values_only=True))))
    assert json.loads(r["crop_box"])["y1"] == 0.6
